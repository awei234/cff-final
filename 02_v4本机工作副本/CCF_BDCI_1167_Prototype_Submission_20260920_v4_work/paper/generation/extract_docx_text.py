"""Extract plain text (paragraphs + tables) from a .docx for faithful reproduction.

Usage: python extract_docx_text.py <input.docx> <output.txt>
Read-only: never writes to the source document.
"""
import sys
from pathlib import Path

from docx import Document


def main() -> int:
    if len(sys.argv) != 3:
        print(__doc__)
        return 2
    src = Path(sys.argv[1])
    dst = Path(sys.argv[2])
    doc = Document(str(src))
    lines = []
    for para in doc.paragraphs:
        lines.append(para.text)
    for ti, table in enumerate(doc.tables):
        lines.append(f"[TABLE {ti}]")
        for row in table.rows:
            lines.append(" | ".join(cell.text.replace("\n", " ") for cell in row.cells))
    dst.write_text("\n".join(lines), encoding="utf-8")
    print("paragraphs:", len(doc.paragraphs), "tables:", len(doc.tables), "->", dst)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
